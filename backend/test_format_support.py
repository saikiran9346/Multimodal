import os
from pathlib import Path

# Create sample files for all 6 new formats
import docx
doc = docx.Document()
doc.add_heading('Technical Specifications', 0)
doc.add_paragraph('Operating temperature range is -20 to +90 degrees Celsius for stainless steel pumps.')
doc.save('sample.docx')

import pptx
prs = pptx.Presentation()
slide = prs.slides.add_slide(prs.slide_layouts[0])
slide.shapes.title.text = 'Pump System Overview'
slide.placeholders[1].text = 'Pumping capacity is 150 m3/h at maximum operating pressure of 16 bar.'
prs.save('sample.pptx')

import openpyxl
wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'Specifications'
ws.append(['Parameter', 'Value', 'Unit'])
ws.append(['Max System Pressure', '16', 'bar'])
ws.append(['Liquid Temperature', '90', 'C'])
wb.save('sample.xlsx')

with open('sample.html', 'w', encoding='utf-8') as f:
    f.write('<!DOCTYPE html><html><body><h1>Maintenance Protocol</h1><p>Inspect motor bearings every 4000 hours of continuous operation.</p></body></html>')

with open('sample.md', 'w', encoding='utf-8') as f:
    f.write('# Installation Guidelines\n\nEnsure foundation weight is at least 1.5 times the pump weight.\n')

with open('sample.csv', 'w', encoding='utf-8') as f:
    f.write('Component,Material,Rating\nShaft Seal,Silicon Carbide,16 bar\nImpeller,Stainless Steel 316,25 bar\n')

from app.ingestion.parser import parse_document
from app.ingestion.chunker import chunk_document
from app.generation.llm import _format_doc_header

sample_files = [
    'sample.docx',
    'sample.pptx',
    'sample.xlsx',
    'sample.html',
    'sample.md',
    'sample.csv'
]

print("\n" + "="*80)
print("MULTI-FORMAT DOCUMENT PARSING & CITATION FALLBACK VERIFICATION TEST")
print("="*80)

for filename in sample_files:
    filepath = Path(filename)
    try:
        doc = parse_document(str(filepath))
        chunks = chunk_document(doc)
        
        for c in chunks:
            c['doc_name'] = filename
            
        page_numbers = [c.get('page_no') for c in chunks]
        sample_headers = [_format_doc_header(c) for c in chunks]
        
        print(f"\n[FORMAT TEST] {filename}")
        print(f"  Parsed Successfully : YES")
        print(f"  Total Chunks        : {len(chunks)}")
        print(f"  Page Numbers        : {page_numbers}")
        print(f"  Formatted Citations : {sample_headers}")
        print(f"  Sample Text Snippet : {chunks[0]['text'][:80] if chunks else 'N/A'}")
    except Exception as e:
        print(f"\n[FORMAT TEST] {filename}")
        print(f"  Parsed Successfully : NO (Error: {str(e)})")

print("\n" + "="*80)
